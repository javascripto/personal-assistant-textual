"""XLSX export for the complete task collection."""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from personal_assistant.models import Task


def export_tasks(tasks: list[Task], directory: Path = Path("exports")) -> Path:
    """Write tasks to a timestamped workbook and return its path."""
    directory.mkdir(parents=True, exist_ok=True)
    output = directory / f"tasks-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Workbook was created without an active worksheet.")
    sheet.title = "Tarefas"
    sheet.append(
        ["Status", "Data", "Tag", "Título", "Descrição", "Concluída em"]
    )

    for task in tasks:
        sheet.append(
            [
                task.status.value,
                task.created_at.strftime("%d/%m/%Y %H:%M"),
                task.tag,
                task.title,
                task.description,
                (
                    task.completed_at.strftime("%d/%m/%Y %H:%M")
                    if task.completed_at
                    else ""
                ),
            ]
        )

    header_fill = PatternFill("solid", fgColor="008888")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for column, width in {
        "A": 14,
        "B": 20,
        "C": 16,
        "D": 40,
        "E": 60,
        "F": 20,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(output)
    return output
