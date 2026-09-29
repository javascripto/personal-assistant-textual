"""Record real terminal interaction and verify its SQLite/XLSX effects."""

from __future__ import annotations

import json
import os
import shutil
import sqlite3
import time
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import TypedDict

import pexpect
import pyte
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "recordings"
TITLE = "Preparar demo do Personal Assistant"
EDITED_TITLE = "Publicar demo do Personal Assistant"
DESCRIPTION = (
    "Objetivo: mostrar um fluxo completo de tarefas no terminal.\n\n"
    "1. Criar uma tarefa com contexto e proximos passos.\n"
    "2. Consultar detalhes e atualizar a descricao.\n"
    "3. Concluir, filtrar e exportar a lista para Excel.\n\n"
    "Tudo fica salvo localmente em SQLite."
)


class Chapter(TypedDict):
    seconds: float
    title: str
    keys: str


class TerminalDemo:
    """Send keys only after reading output, with verifiable checkpoints."""

    def __init__(self, workspace: Path, recording: Path) -> None:
        self.workspace = workspace
        self.started = time.monotonic()
        self.chapters: list[Chapter] = []
        self.checks: list[str] = []
        self.screen = pyte.Screen(120, 36)
        self.stream = pyte.Stream(self.screen)
        self.transcript = ""
        env = dict(os.environ)
        for name in ("NO_COLOR", "FORCE_COLOR", "ASCIINEMA_SESSION"):
            env.pop(name, None)
        env.update(
            TERM="xterm-256color",
            COLORTERM="truecolor",
            PS1="demo $ ",
            PROMPT_COMMAND="",
            BASH_SILENCE_DEPRECATION_WARNING="1",
            UV_NO_SYNC="1",
            UV_PROJECT_ENVIRONMENT=str(ROOT / ".venv"),
        )
        self.child: pexpect.spawn[str] = pexpect.spawn(
            "asciinema",
            [
                "rec",
                "--quiet",
                "--return",
                "--output-format",
                "asciicast-v2",
                "--window-size",
                "120x36",
                "--title",
                "Personal Assistant | demo verificada",
                "--command",
                "/bin/bash --noprofile --norc",
                str(recording),
            ],
            cwd=str(workspace),
            env=env,
            encoding="utf-8",
            dimensions=(36, 120),
            timeout=60,
        )

    def pump(self, seconds: float = 0.4) -> None:
        """Continuously drain the PTY so the child never blocks on output."""
        deadline = time.monotonic() + seconds
        while time.monotonic() < deadline:
            try:
                output = self.child.read_nonblocking(
                    65536, timeout=min(0.1, deadline - time.monotonic())
                )
            except pexpect.TIMEOUT:
                continue
            except pexpect.EOF:
                return
            self.transcript += output
            self.stream.feed(output)

    def visible(self) -> str:
        return "\n".join(self.screen.display)

    def wait_for(self, text: str, timeout: float = 15) -> None:
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            self.pump(0.1)
            if text in self.visible():
                return
        raise AssertionError(f"Tela não encontrada: {text}\n{self.visible()}")

    def key(self, value: str, pause: float = 0.5) -> None:
        self.child.send(value)
        self.pump(pause)

    def type_text(self, value: str) -> None:
        for char in value:
            self.key(char, 0.035)

    def paste(self, value: str) -> None:
        self.key("\x1b[200~" + value + "\x1b[201~", 0.7)

    def chapter(self, title: str, keys: str = "") -> None:
        self.chapters.append(
            {
                "seconds": round(time.monotonic() - self.started, 2),
                "title": title,
                "keys": keys,
            }
        )
        print(f"Demo: {title}", flush=True)
        self.pump(2)

    def command(self, command: str, timeout: float = 90) -> None:
        self.key("\x0c")
        self.type_text(command + " || exit 1")
        self.key("\r")
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            self.pump(0.1)
            if self.screen.display[self.screen.cursor.y].strip() == "demo $":
                break
        else:
            raise AssertionError(f"Comando não terminou: {command}")
        self.pump(2)

    def sql(self, query: str, parameters: tuple[str, ...] = ()) -> str:
        with sqlite3.connect(
            self.workspace / "data" / "personal_assistant.db"
        ) as database:
            row = database.execute(query, parameters).fetchone()
            return "" if row is None else str(row[0])

    def check(self, name: str, condition: bool) -> None:
        if not condition:
            raise AssertionError(f"{name}\n{self.visible()}")
        self.checks.append(name)

    def run(self) -> None:
        self.wait_for("$ ")
        self.key("export PS1='demo $ '\r")
        self.wait_for("demo $ ")
        self.key("\x0c")
        self.chapter("Personal Assistant", "Python + Textual + SQLite")
        self.type_text("make run")
        self.key("\r")
        self.wait_for("Ordenado por:")
        self.chapter("Sua lista, direto no terminal", "Setas: navegar")
        self.check(
            "seed inicial: 10 tarefas",
            self.sql("SELECT COUNT(*) FROM tasks") == "10",
        )
        self.key("\x1b[B", 1)
        self.key("\x1b[B", 1)
        self.key("\x1b[A", 1)

        self.chapter("Criar tarefa com descrição longa", "N | Tab | Confirmar")
        self.key("n")
        self.wait_for("Nova tarefa")
        self.type_text(TITLE)
        self.key("\t")
        self.key("\x01")
        self.key("\x0b")
        self.type_text("demo")
        self.key("\t")
        self.paste(DESCRIPTION)
        self.pump(3)
        self.key("\t", 1)
        self.key("\r", 1)
        self.check(
            "criacao confirmada com descricao",
            self.sql("SELECT description FROM tasks WHERE title = ?", (TITLE,))
            == DESCRIPTION,
        )
        # DataTable Home is horizontal; Ctrl+Home selects the first row.
        self.key("\x1b[1;5H")
        self.key("\r")
        self.wait_for("Detalhes da tarefa")
        self.check("detalhes da tarefa criada", TITLE in self.visible())
        self.chapter(
            "Contexto e próximos passos", "Enter: detalhes | Esc: voltar"
        )
        self.pump(3)

        self.chapter("Editar pelos detalhes", "E | Tab | Confirmar")
        self.key("e")
        self.wait_for("Editar tarefa")
        self.key("\x01")
        self.key("\x0b")
        self.type_text(EDITED_TITLE)
        self.key("\t")
        self.key("\t")
        self.key("\x05")
        self.paste("\n\nRevisao: incluir exportacao XLSX e temas.")
        self.key("\t", 1)
        self.key("\r", 1)
        self.check(
            "edicao persistida",
            "Revisao:"
            in self.sql(
                "SELECT description FROM tasks WHERE title = ?", (EDITED_TITLE,)
            ),
        )
        self.key("\r")
        self.wait_for("Detalhes da tarefa")
        self.check("detalhes atualizados", EDITED_TITLE in self.visible())
        self.pump(3)
        self.key("\x1b")

        self.chapter("Concluir e reabrir sem perder a seleção", "Space")
        for expected in ("completed", "pending", "completed"):
            self.key(" ", 1.5)
            self.check(
                f"status {expected}",
                self.sql(
                    "SELECT status FROM tasks WHERE title = ?", (EDITED_TITLE,)
                )
                == expected,
            )

        self.chapter("Filtros por status", "Tab | Setas | Tab")
        self.key("\t", 1)
        self.key("\x1b[C", 2)
        self.key("\x1b[C", 2)
        self.check("filtro concluidos", EDITED_TITLE in self.visible())
        self.key("\x1b[D", 2)
        self.check("filtro pendentes", EDITED_TITLE not in self.visible())
        self.key("\x1b[D", 2)
        self.key("\t")

        self.chapter("Ordenação por data", "O: crescente / decrescente")
        self.key("o", 2)
        self.check(
            "ordenacao crescente",
            self.sql("SELECT value FROM settings WHERE key = 'task_order'")
            == "asc",
        )
        self.key("o", 2)

        self.chapter("Temas para toda a interface", "T: próximo tema")
        self.key("t", 3)
        self.key("t", 3)
        self.check(
            "tema persistido",
            self.sql("SELECT value FROM settings WHERE key = 'theme'")
            != "personal-assistant-retro",
        )
        self.chapter("Command palette", "Ctrl+P | Select theme")
        self.key("\x10", 1)
        self.type_text("theme")
        self.pump(2)
        self.key("\r", 1)
        self.type_text("personal-assistant-retro")
        self.pump(2)
        self.key("\r", 2)
        self.check(
            "retorno ao tema retro via palette",
            self.sql("SELECT value FROM settings WHERE key = 'theme'")
            == "personal-assistant-retro",
        )

        self.chapter("Exportar todas as tarefas", "X: planilha XLSX")
        self.key("x", 3)
        exports = list((self.workspace / "exports").glob("*.xlsx"))
        self.check("arquivo XLSX criado", len(exports) == 1)
        workbook = load_workbook(exports[0], read_only=True)
        try:
            sheet = workbook.active
            assert sheet is not None
            self.check(
                "XLSX inclui titulo e descricao",
                any(
                    row[3] == EDITED_TITLE and "Revisao:" in str(row[4])
                    for row in sheet.iter_rows(min_row=2, values_only=True)
                ),
            )
        finally:
            workbook.close()

        self.chapter("Cancelar sem alterar dados", "N | Esc; E | Esc")
        self.key("n")
        self.type_text("Esta tarefa sera cancelada")
        self.key("\x1b", 1)
        self.check(
            "criacao cancelada", self.sql("SELECT COUNT(*) FROM tasks") == "11"
        )
        self.key("e")
        self.type_text("Rascunho")
        self.key("\x1b", 1)
        self.check(
            "edicao cancelada",
            self.sql(
                "SELECT COUNT(*) FROM tasks WHERE title = ?", (EDITED_TITLE,)
            )
            == "1",
        )

        self.chapter("Excluir com confirmação", "D | Cancelar / Excluir")
        self.key("d", 2)
        self.wait_for("Excluir tarefa?")
        self.key("\x1b", 1)
        self.check(
            "exclusao cancelada", self.sql("SELECT COUNT(*) FROM tasks") == "11"
        )
        self.key("d", 2)
        self.key("\r", 1)
        self.check(
            "exclusao confirmada",
            self.sql(
                "SELECT COUNT(*) FROM tasks WHERE title = ?", (EDITED_TITLE,)
            )
            == "0",
        )

        self.chapter("Persistência entre execuções", "Q | make run")
        self.key("q")
        self.wait_for("demo $ ")
        self.type_text("make run")
        self.key("\r")
        self.wait_for("Ordenado por:")
        self.pump(3)
        self.key("q")
        self.wait_for("demo $ ")

        self.chapter("Workflow de desenvolvimento", "Makefile + uv")
        self.command("make help")
        self.command("make format")
        self.command("make check")
        self.check("gate de qualidade", "passed" in self.transcript)
        self.chapter("Seed idempotente com backup", "make seed")
        self.command("make seed")
        self.check(
            "backup antes do seed",
            bool(list((self.workspace / "data").glob("*.backup-*.db"))),
        )
        self.check(
            "seed sem duplicacao",
            self.sql("SELECT COUNT(*) FROM tasks") == "10",
        )
        self.chapter("Limpar caches e preservar dados", "make clear")
        self.command("make clear")
        self.check(
            "banco e exportacao preservados",
            exports[0].exists()
            and self.sql("SELECT COUNT(*) FROM tasks") == "10",
        )
        self.command("ls data/ exports/")
        self.chapter("Personal Assistant", "Tarefas locais. Fluxo por teclado.")
        self.key("exit\r")
        self.child.expect(pexpect.EOF, timeout=10)
        self.child.close()
        self.check("sessao encerrou com sucesso", self.child.exitstatus == 0)
        self.check("sem excecoes na TUI", "Traceback" not in self.transcript)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix="personal-assistant-demo-") as directory:
        workspace = Path(directory)
        for name in (
            "Makefile",
            "pyproject.toml",
            "uv.lock",
            "README.md",
            "pyrightconfig.json",
        ):
            shutil.copy2(ROOT / name, workspace / name)
        for name in ("src", "tests"):
            shutil.copytree(
                ROOT / name,
                workspace / name,
                ignore=shutil.ignore_patterns("__pycache__"),
            )
        (workspace / ".venv").symlink_to(
            ROOT / ".venv", target_is_directory=True
        )
        recording = workspace / "demo.cast"
        demo = TerminalDemo(workspace, recording)
        try:
            demo.run()
            for line in recording.read_text().splitlines():
                json.loads(line)
            shutil.copy2(recording, OUTPUT / "personal-assistant-demo.cast")
            (OUTPUT / "demo-chapters.json").write_text(
                json.dumps(demo.chapters, ensure_ascii=False, indent=2) + "\n"
            )
            (OUTPUT / "demo-validation.json").write_text(
                json.dumps({"checks": demo.checks}, indent=2) + "\n"
            )
            print(f"Gravação validada: {len(demo.checks)} verificações.")
        finally:
            demo.child.close(force=True)


if __name__ == "__main__":
    main()
